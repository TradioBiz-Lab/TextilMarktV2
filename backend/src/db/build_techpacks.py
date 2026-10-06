#!/usr/bin/env python3
"""Builds the demo tech packs (PDF) and measurement sheets (XLSX) for every style in demoStyles.js.

Each one is a MODIFIED version of the matching tech pack from the "Tradio Costing" folder: same sheet
structure (style sheet, construction, finishing, print/logo, stitch details, spec, how to measure, colour
ways), but rebranded to Aero Active with our style codes, a different designer/collection, adjusted fabric
GSM and measurements, and our own flat sketches (derived from the product images, not the originals).

Usage (from this folder):  python3 build_techpacks.py "/path/to/Tradio Costing"
Output: demo-docs/techpack-<slug>.pdf and demo-docs/measurements-<slug>.xlsx
"""
import json, os, re, subprocess, sys
from PIL import Image, ImageDraw, ImageFont, ImageFilter, ImageChops, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from tp_extract import extract

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'demo-docs'); os.makedirs(OUT, exist_ok=True)
IMG = os.path.join(HERE, 'demo-images')
SRC_ROOT = sys.argv[1] if len(sys.argv) > 1 else '/Users/ankitbera/Documents/Claude/Projects/Tradio Consolidated/Tradio Costing'

W, H = 1169, 827
INK, ACCENT, GREY, SOFT = (34, 36, 42), (234, 88, 12), (120, 124, 132), (250, 236, 226)
F = '/System/Library/Fonts/Supplemental/'
def font(sz, bold=False): return ImageFont.truetype(F + ('Arial Bold.ttf' if bold else 'Arial.ttf'), sz)

COLOURS = {'Olive Green': '#6B7A4A', 'Black': '#1c1d21', 'Teal': '#0f766e', 'Royal Blue': '#2a4fb8', 'Red': '#c0392b', 'Black Floral': '#2b2d33',
           'Charcoal': '#4b5058', 'Orange Red': '#e8501c', 'White': '#f4f4f2', 'Navy': '#1b2a4e', 'Orange': '#e8741f'}
STITCHES = ['Three Thread Overlock', 'Four Thread Overlock', 'Two Needle Flatlock Stitch', '3 Needle 5 Thread Flatseam', 'Two Needle Cover Stitch',
            'One Needle Lock Stitch', 'Satin Stitch', 'Bartack Stitch', 'Zig Zag Stitch', 'Free Cut Hem', 'Laser Cut']

def rebrand(t):
    t = t.replace('\u2014', '-').replace('\u2013', '-')   # no em/en dashes in our documents
    t = re.sub(r'fitleasure', 'Aero Active', t, flags=re.I)
    t = re.sub(r'\bSiver\b', 'Silver', t)
    return re.sub(r'\s+', ' ', t).strip()

# ───────────── header values (modified from the source) ─────────────
def mod_fabric(s):
    m = re.match(r'(\d+)% (\w+) (\d+)% (\w+)', s or '')
    if m:  # shift the blend by two points
        a, b = int(m.group(1)) - 2, int(m.group(3)) + 2
        return f'{a}% {m.group(2)} {b}% {m.group(4)}'
    return '88% Polyester 12% Spandex' if (not s or 'Blend' in s or '%' not in s) else s.replace('Spendex', 'Spandex')
def mod_gsm(s):
    nums = re.findall(r'\d+', s or '')
    if len(nums) == 2: return f'{int(nums[0]) + 10} - {int(nums[1]) + 10}'
    if len(nums) == 1: return str(int(nums[0]) + 10)
    return s or ''
def header_for(style, src):
    h = src['header']
    return [('Brand', 'Aero Active'), ('Style No.', style['style']), ('Description', style['product']), ('Category', h.get('Category', '')), ('Gender', style.get('gender') or h.get('Gender', ''))], \
           [('Designer', 'Aero Design Team'), ('Fabric Comp.', mod_fabric(h.get('Fabric Comp.'))), ('GSM', mod_gsm(h.get('GSM'))), ('Size Range', h.get('Size Range', 'S-XXL')), ('Collection', 'SS27 Core Essentials')]

# ───────────── drawing helpers ─────────────
def dashed_line(d, p1, p2, color=ACCENT, width=2, dash=7, gap=5):
    (x1, y1), (x2, y2) = p1, p2
    L = max(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5, 1); ux, uy = (x2 - x1) / L, (y2 - y1) / L
    s = 0
    while s < L:
        e = min(s + dash, L); d.line([(x1 + ux * s, y1 + uy * s), (x1 + ux * e, y1 + uy * e)], fill=color, width=width); s += dash + gap
def dashed_rect(d, box, color=ACCENT, width=2):
    x1, y1, x2, y2 = box
    for a, b in [((x1, y1), (x2, y1)), ((x2, y1), (x2, y2)), ((x2, y2), (x1, y2)), ((x1, y2), (x1, y1))]: dashed_line(d, a, b, color, width)
def wrap(d, text, fnt, maxw):
    words, lines, cur = text.split(), [], ''
    for w in words:
        t = (cur + ' ' + w).strip()
        if d.textlength(t, font=fnt) <= maxw: cur = t
        else: lines.append(cur); cur = w
    if cur: lines.append(cur)
    return lines

def frame(style, src, page_no, sheet):
    im = Image.new('RGB', (W, H), 'white'); d = ImageDraw.Draw(im)
    dashed_rect(d, (20, 40, W - 20, H - 32))
    d.text((W - 24, 12), 'Drop | SS27', font=font(15), fill=ACCENT, anchor='ra')
    # logo cell
    d.line([(146, 41), (146, 170)], fill=ACCENT, width=1)
    d.polygon([(66, 128), (90, 78), (114, 128), (102, 128), (90, 100), (78, 128)], fill=INK)
    d.text((90, 144), 'AERO ACTIVE', font=font(13, True), fill=INK, anchor='ma')
    left, right = header_for(style, src)
    for i, (k, v) in enumerate(left):
        y = 55 + i * 24.5; d.text((160, y), k, font=font(15, True), fill=INK); d.text((312, y), v, font=font(15), fill=INK)
    d.line([(636, 41), (636, 170)], fill=ACCENT, width=1)
    for i, (k, v) in enumerate(right):
        y = 55 + i * 24.5; d.text((648, y), k, font=font(15, True), fill=INK); d.text((795, y), v, font=font(15), fill=INK)
    d.line([(21, 170), (W - 21, 170)], fill=ACCENT, width=1)
    d.text((22, H - 24), f'Sheet Name : {sheet}', font=font(16), fill=ACCENT)
    d.text((W - 24, H - 24), f'Page {page_no}', font=font(16), fill=ACCENT, anchor='ra')
    return im, d

# ───────────── flat sketches made from the product image ─────────────
def sketch(slug, height, back=False):
    """A white-filled, dark-outlined flat of the garment (smooth silhouette plus interior seam lines),
    derived from the product image with edge detection."""
    import cv2, numpy as np
    bgr = cv2.imread(os.path.join(IMG, slug + '.jpg')); gray = cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY)
    # light garments on a white background (white tee, white tank) are barely darker than the backdrop,
    # so use a sensitive threshold and a wide close; the largest connected blob is the garment.
    mask = (gray < 247).astype(np.uint8) * 255
    k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (23, 23)); mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, k)
    n, lab, stats, _ = cv2.connectedComponentsWithStats(mask)
    if n > 1: mask = (lab == 1 + int(np.argmax(stats[1:, cv2.CC_STAT_AREA]))).astype(np.uint8) * 255
    ff = mask.copy(); cv2.floodFill(ff, np.zeros((mask.shape[0] + 2, mask.shape[1] + 2), np.uint8), (0, 0), 255)
    mask = cv2.bitwise_or(mask, cv2.bitwise_not(ff))                       # fill interior holes
    mask = (cv2.GaussianBlur(mask, (0, 0), 2.2) > 127).astype(np.uint8) * 255  # smooth the outline
    x, y, w, h = cv2.boundingRect(mask); pad = 12
    x0, y0, x1, y1 = max(x - pad, 0), max(y - pad, 0), min(x + w + pad, gray.shape[1]), min(y + h + pad, gray.shape[0])
    gray, mask = gray[y0:y1, x0:x1], mask[y0:y1, x0:x1]
    sc = height / gray.shape[0]; size = (int(gray.shape[1] * sc), height)
    gray = cv2.resize(gray, size, interpolation=cv2.INTER_CUBIC); mask = (cv2.resize(mask, size, interpolation=cv2.INTER_LINEAR) > 127).astype(np.uint8) * 255
    out = np.zeros((height, size[0], 4), np.uint8); out[mask > 0] = (255, 255, 255, 255)
    if not back:
        e = cv2.Canny(cv2.GaussianBlur(gray, (0, 0), 1.6), 18, 52)
        inside = cv2.erode(mask, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (15, 15)))
        e = cv2.bitwise_and(e, inside); e = cv2.dilate(e, np.ones((2, 2), np.uint8))
        n, lab, stats, _ = cv2.connectedComponentsWithStats(e)
        keep = np.isin(lab, [i for i in range(1, n) if stats[i, cv2.CC_STAT_AREA] >= 90])
        out[keep] = (125, 129, 138, 255)
    cnts, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_NONE)
    cv2.polylines(out, cnts, True, INK + (255,), 3, cv2.LINE_AA)
    img = Image.fromarray(out)
    if back: ImageDraw.Draw(img).line([(img.width // 2, int(img.height * 0.2)), (img.width // 2, int(img.height * 0.82))], fill=(125, 129, 138, 255), width=2)
    return img.transpose(Image.FLIP_LEFT_RIGHT) if back else img

def paste_center(im, sk, cx, top):
    im.paste(sk, (int(cx - sk.width / 2), top), sk); return (int(cx - sk.width / 2), top, int(cx + sk.width / 2), top + sk.height)

def callout_page(style, src, n, sheet, items):
    im, d = frame(style, src, n, sheet)
    items = [rebrand(c['text']) for c in items if 'Reference' not in c['text']][:8]
    if not items: items = ['Seams finished with overlock', 'Hem finished with flatlock stitch']
    sk = sketch(style['slug'], 500); box = paste_center(im, sk, W / 2, 215)
    d = ImageDraw.Draw(im)
    left, right = items[0::2], items[1::2]
    span = box[3] - box[1]
    for side, lst in (('L', left), ('R', right)):
        for i, text in enumerate(lst):
            yc = box[1] + span * (i + 0.5) / max(len(lst), 1)
            lines = wrap(d, text, font(15), 250); y0 = yc - 9 * len(lines)
            if side == 'L':
                for j, ln in enumerate(lines): d.text((300, y0 + j * 19), ln, font=font(15), fill=INK, anchor='ra')
                dashed_line(d, (308, yc), (box[0] + (box[2] - box[0]) * 0.34, yc))
            else:
                for j, ln in enumerate(lines): d.text((W - 300, y0 + j * 19), ln, font=font(15), fill=INK)
                dashed_line(d, (W - 308, yc), (box[2] - (box[2] - box[0]) * 0.34, yc))
    return im

def style_page(style, src):
    im, d = frame(style, src, 1, 'Style Sheet')
    sk = sketch(style['slug'], 440); sb = sketch(style['slug'], 440, back=True)
    fit = min(1.0, 740 / (sk.width + sb.width + 40))
    if fit < 1: sk = sk.resize((int(sk.width * fit), int(sk.height * fit)), Image.LANCZOS); sb = sb.resize((int(sb.width * fit), int(sb.height * fit)), Image.LANCZOS)
    cxf = 40 + sk.width / 2; cxb = 40 + sk.width + 40 + sb.width / 2
    b1 = paste_center(im, sk, cxf, 440 - sk.height + 225); b2 = paste_center(im, sb, cxb, 440 - sb.height + 225); d = ImageDraw.Draw(im)
    d.text(((b1[0] + b1[2]) / 2, 700), 'FRONT', font=font(18, True), fill=(220, 38, 38), anchor='ma')
    d.text(((b2[0] + b2[2]) / 2, 700), 'BACK', font=font(18, True), fill=(220, 38, 38), anchor='ma')
    ref = Image.open(os.path.join(IMG, style['slug'] + '.jpg')).convert('RGB').resize((250, 250), Image.LANCZOS)
    im.paste(ref, (860, 290)); d.rectangle([860, 290, 1110, 540], outline=(190, 190, 194), width=1)
    d.text((985, 560), 'Product Reference Image', font=font(18, True), fill=(220, 38, 38), anchor='ma')
    return im

def print_page(style, src, n):
    im, d = frame(style, src, n, 'Logo CAD')
    txt = list(dict.fromkeys(rebrand(c['text']) for c in src.get('print', [])))
    tech = [t for t in txt if re.match(r'(Print Technique|Temperature|Time|Print Colour)', t)] or ['Print Technique : Silver Reflective Heatset', 'Temperature : 150°C - 160°C', 'Time : 8-10 secs']
    dims = [t for t in txt if re.match(r'^[\d.]+ ?cms?$', t)][:3] or ['4.5 cm', '0.5 cm']
    place = [t for t in txt if re.search(r'(aligned|above|from|below)', t, re.I) and not re.match(r'(Print|Temp|Time)', t)][:3] or ['Aligned to center back']
    sk = sketch(style['slug'], 340); box = paste_center(im, sk, 330, 300); d = ImageDraw.Draw(im)
    lx, ly = int((box[0] + box[2]) / 2) - 55, box[1] + 90
    d.rectangle([lx, ly, lx + 110, ly + 22], outline=ACCENT, width=2); d.text((lx + 55, ly + 4), 'AERO ACTIVE', font=font(12, True), fill=INK, anchor='ma')
    d.line([(lx, ly + 34), (lx + 110, ly + 34)], fill=ACCENT, width=2); d.text((lx + 55, ly + 38), dims[0], font=font(14, True), fill=ACCENT, anchor='ma')
    d.text((330, 250), 'LOGO placement (to scale)', font=font(16, True), fill=INK, anchor='ma')
    y = 290
    for t in tech:
        d.text((760, y), t, font=font(17), fill=INK); y += 30
    y += 14; d.text((760, y), 'Placement', font=font(17, True), fill=INK); y += 30
    for t in place:
        for ln in wrap(d, t, font(16), 340): d.text((760, y), ln, font=font(16), fill=INK); y += 22
        y += 6
    dashed_rect(d, (740, 270, 1120, y + 12), width=1)
    return im

def stitch_page(style, src, n):
    im, d = frame(style, src, n, 'Stitch Details')
    d.text((60, 205), 'STITCH LEGEND :', font=font(17, True), fill=INK)
    for i, s in enumerate(STITCHES):
        y = 245 + i * 38; d.ellipse([60, y, 90, y + 30], outline=ACCENT, width=2); d.text((75, y + 15), str(i + 1), font=font(15, True), fill=ACCENT, anchor='mm'); d.text((108, y + 6), s, font=font(17), fill=INK)
    sk = sketch(style['slug'], 500); box = paste_center(im, sk, 760, 215); d = ImageDraw.Draw(im)
    pts = [(0.5, 0.08, 2), (0.3, 0.28, 1), (0.7, 0.28, 1), (0.25, 0.55, 5), (0.75, 0.55, 5), (0.5, 0.6, 6), (0.5, 0.92, 6), (0.35, 0.9, 3), (0.65, 0.9, 3)]
    for fx, fy, k in pts:
        x = box[0] + (box[2] - box[0]) * fx; y = box[1] + (box[3] - box[1]) * fy
        d.ellipse([x - 11, y - 11, x + 11, y + 11], fill='white', outline=ACCENT, width=2); d.text((x, y), str(k), font=font(13, True), fill=ACCENT, anchor='mm')
    return im

def mod_val(v): return round(v * 1.02 * 2) / 2 if v >= 10 else v
def spec_page(style, src, n):
    im, d = frame(style, src, n, 'Spec Sheet')
    sp = src.get('spec', {'sizes': ['S', 'M', 'L', 'XL', 'XXL'], 'rows': []}); sizes = sp['sizes']
    rows = sp['rows'][:20]; rh = 27 if len(rows) <= 17 else 24
    x0, ytop = 60, 190
    cols = [60, 120, 520, 600] + [600 + 70 * (i + 1) for i in range(len(sizes))]
    d.rectangle([x0, ytop, cols[-1], ytop + rh], fill=SOFT)
    for t, x in zip(['Code', 'Measurement point', 'Tol'] + sizes, [66, 126, 530] + [606 + 70 * (i + 1) for i in range(len(sizes))]):
        d.text((x, ytop + 5), t, font=font(15, True), fill=INK)
    for i, r in enumerate(rows):
        y = ytop + rh * (i + 1); d.line([(x0, y), (cols[-1], y)], fill=(225, 225, 228), width=1)
        d.text((66, y + 5), r['code'], font=font(15, True), fill=INK); d.text((126, y + 5), rebrand(r['name'])[:50], font=font(15), fill=INK); d.text((530, y + 5), r['tol'], font=font(15), fill=GREY)
        for j, v in enumerate(r['vals'][:len(sizes)]):
            val = mod_val(v); d.text((606 + 70 * (j + 1), y + 5), (f'{val:g}'), font=font(15), fill=INK)
    d.text((60, H - 56), 'All measurements in cm, garment laid flat. Tolerances apply to finished garment.', font=font(13), fill=GREY)
    return im, [dict(r, vals=[mod_val(v) for v in r['vals']]) for r in sp['rows']], sizes

def measure_page(style, src, n, rows):
    im, d = frame(style, src, n, 'How To Measure')
    sk = sketch(style['slug'], 500); box = paste_center(im, sk, 330, 215); d = ImageDraw.Draw(im)
    letters = [r['code'] for r in rows[:14]] or list('ABCDEFGH')
    import math
    cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2; rx, ry = (box[2] - box[0]) * 0.34, (box[3] - box[1]) * 0.4
    for i, L in enumerate(letters):
        a = -math.pi / 2 + 2 * math.pi * i / len(letters); x, y = cx + rx * math.cos(a), cy + ry * math.sin(a)
        d.ellipse([x - 12, y - 12, x + 12, y + 12], fill='white', outline=ACCENT, width=2); d.text((x, y), L, font=font(14, True), fill=ACCENT, anchor='mm')
    y = 205
    for r in rows[:14]:
        d.text((700, y), r['code'], font=font(15, True), fill=ACCENT); d.text((735, y), rebrand(r['name'])[:48], font=font(15), fill=INK); y += 28
    return im

def colour_page(style, src, n):
    im, d = frame(style, src, n, 'Colour Ways')
    x = 90
    for i, c in enumerate(style['colours']):
        hx = COLOURS.get(c, '#888888'); d.rectangle([x, 230, x + 220, 450], fill=hx, outline=(190, 190, 194), width=1)
        d.text((x, 465), c, font=font(20, True), fill=INK); d.text((x, 495), f'Colourway {i + 1}  |  Ref AER-C{i + 1:02d}', font=font(15), fill=GREY); x += 260
    d.text((90, 600), 'Colours to be matched to the approved lab dip. Bulk shade tolerance: Delta E 1.0 (CMC 2:1).', font=font(15), fill=GREY)
    return im

def xlsx(style, src, rows, sizes, path):
    wb = Workbook(); ws = wb.active; ws.title = 'Measurement Spec'
    thin = Side(style='thin', color='DDDDDD'); head = PatternFill('solid', fgColor='FFEDD5')
    ws['A1'] = f"Measurement Spec - {style['product']}"; ws['A1'].font = Font(bold=True, size=14)
    meta = [('Brand', 'Aero Active'), ('Style No.', style['style']), ('Fabric', mod_fabric(src['header'].get('Fabric Comp.'))), ('GSM', mod_gsm(src['header'].get('GSM'))), ('Size range', src['header'].get('Size Range', 'S-XXL')), ('Unit', 'cm, garment laid flat')]
    for i, (k, v) in enumerate(meta): ws.cell(row=2 + i, column=1, value=k).font = Font(bold=True); ws.cell(row=2 + i, column=2, value=v)
    r0 = 9
    for j, h in enumerate(['Code', 'Measurement point', 'Tol (cm)'] + sizes):
        c = ws.cell(row=r0, column=1 + j, value=h); c.font = Font(bold=True); c.fill = head; c.border = Border(bottom=thin); c.alignment = Alignment(horizontal='center' if j != 1 else 'left')
    for i, r in enumerate(rows):
        ws.cell(row=r0 + 1 + i, column=1, value=r['code']); ws.cell(row=r0 + 1 + i, column=2, value=rebrand(r['name'])); ws.cell(row=r0 + 1 + i, column=3, value=r['tol'])
        for j, v in enumerate(r['vals'][:len(sizes)]): c = ws.cell(row=r0 + 1 + i, column=4 + j, value=v); c.alignment = Alignment(horizontal='center')
    ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 46; ws.column_dimensions['C'].width = 10
    for j in range(len(sizes)): ws.column_dimensions[chr(ord('D') + j)].width = 9
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    n = wb.create_sheet('Notes')
    for i, t in enumerate(['Measure the garment flat, relaxed, on a table. Half-width measurements are taken straight across.', 'Tolerances apply to the finished garment after wash.', 'Grade between sizes as per the table; do not interpolate by eye.', 'Any measurement outside tolerance is a Major defect at final inspection.']):
        n.cell(row=1 + i, column=1, value=t)
    n.column_dimensions['A'].width = 100
    wb.save(path)

def build(style):
    src = extract(os.path.join(SRC_ROOT, style['source']))
    cons = src.get('construction', []); fin = src.get('finishing') or cons[len(cons) // 2:]
    pages = [style_page(style, src), callout_page(style, src, 2, 'Construction Sheet', cons), callout_page(style, src, 3, 'Finishing Details', fin), print_page(style, src, 4), stitch_page(style, src, 5)]
    sp, rows, sizes = spec_page(style, src, 6)
    pages += [sp, measure_page(style, src, 7, rows), colour_page(style, src, 8)]
    pdf = os.path.join(OUT, f"techpack-{style['slug']}.pdf")
    pages[0].save(pdf, 'PDF', save_all=True, append_images=pages[1:], resolution=100.0, quality=72)
    xlsx(style, src, rows, sizes, os.path.join(OUT, f"measurements-{style['slug']}.xlsx"))
    # Fabric facts as printed on this tech pack, so the order page can show the same numbers.
    mp = os.path.join(OUT, 'fabric-meta.json')
    meta = json.load(open(mp)) if os.path.exists(mp) else {}
    meta[style['slug']] = dict(composition=mod_fabric(src['header'].get('Fabric Comp.')), gsm=mod_gsm(src['header'].get('GSM')))
    json.dump(meta, open(mp, 'w'), indent=1)
    return pdf

if __name__ == '__main__':
    styles = json.loads(subprocess.run(['node', '-e', "import('./demoStyles.js').then(m=>console.log(JSON.stringify(m.STYLES)))"], capture_output=True, text=True, cwd=HERE).stdout)
    only = os.environ.get('ONLY')
    for st in styles:
        if only and st['slug'] != only: continue
        p = build(st); print(f"{st['slug']:26} {os.path.getsize(p) // 1024:5d} KB  <- {os.path.basename(st['source'])}")
